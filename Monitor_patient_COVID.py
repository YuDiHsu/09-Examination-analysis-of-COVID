import openpyxl
import os
import glob
import re
import xlsxwriter
import pandas as pd
import requests
import cx_Oracle
import datetime
from smtp import SMTP
from subprocess import Popen
import time


def data_clean():
    # 資料清理
    report_number_list = []
    data_list = []

    _data = pd.read_csv(os.path.join('.', 'imported_data', 'cdc', '2019-nCov_all_exam.csv'), usecols=['NAME', 'SAMPLE', 'REPORT', 'SAMPLE_DATE',
                                                                                                      'RESULT_DESC', 'SAMPLE_TYPE_DESC'], index_col=False)

    raw_data = _data[(_data['SAMPLE_TYPE_DESC'] == '下呼吸道抽取物') | (_data['SAMPLE_TYPE_DESC'] == '痰')
                     | (_data['SAMPLE_TYPE_DESC'] == '痰液') | (_data['SAMPLE_TYPE_DESC'] == '鼻咽拭子/咽喉擦拭-病毒') | (_data['SAMPLE_TYPE_DESC'] == '咽喉擦拭液')
                     | (_data['SAMPLE_TYPE_DESC'] == '鼻咽擦拭液') | (_data['SAMPLE_TYPE_DESC'] == '唾液(機場專用)')]

    raw_data['SAMPLE_TYPE_DESC'] = raw_data['SAMPLE_TYPE_DESC'].apply(lambda x: x.replace('痰', '痰液') if x == '痰' else x). \
        apply(lambda x: x.replace('鼻咽拭子/咽喉擦拭-病毒', '咽喉擦拭液/鼻咽拭子') if x == '鼻咽拭子/咽喉擦拭-病毒' else x). \
        apply(lambda x: x.replace('咽喉擦拭液', '咽喉擦拭液/鼻咽拭子') if x == '咽喉擦拭液' else x). \
        apply(lambda x: x.replace('鼻咽擦拭液', '咽喉擦拭液/鼻咽拭子') if x == '鼻咽擦拭液' else x)
    raw_data['RESULT_DESC'] = raw_data['RESULT_DESC'].fillna('')

    result_data = pd.read_csv(os.path.join('.', 'imported_data', 'cdc', 'asd.csv'), usecols=['SAMPLE', 'DISEASE', 'COMMENTS'], index_col=False)
    result_data['COMMENTS'] = result_data['COMMENTS'].fillna('')

    case_n_data = pd.read_excel(
        os.path.join('.', 'imported_data', 'cdc', '法傳編號與案次號對照表.xlsx'), header=[1], usecols=['傳染病報告單電腦編號', '案號'], index_col=False,
    ).rename(columns={'傳染病報告單電腦編號': 'REPORT', '案號': 'case_n'})

    report_number_list = list(case_n_data.loc[:, 'REPORT'].unique())

    m_data = pd.merge(raw_data, result_data, on=['SAMPLE'], how='inner')
    m_data['REPORT'] = m_data['REPORT'].astype(int)
    final_data = pd.merge(case_n_data, m_data, on=['REPORT'], how='inner')
    final_data = final_data.loc[:, ~final_data.columns.str.contains('^Unnamed')].drop(['DISEASE'], axis=1)
    final_data['COMMENTS'] = final_data['COMMENTS'].fillna('')
    final_data['COMMENTS'] = final_data['COMMENTS'].apply(lambda x: '無資料' if not x else x)
    final_data['SAMPLE_DATE'] = final_data['SAMPLE_DATE'].fillna('')

    # ['REPORT', 'case_n', 'NAME', 'SAMPLE', 'SAMPLE_DATE', 'SAMPLE_TYPE_DESC', 'RESULT_DESC', 'COMMENTS']
    for row in final_data.itertuples():
        for case in report_number_list:
            if case == row.REPORT:
                # 小於等於今日的資料才要
                if datetime.datetime.strptime(row.SAMPLE_DATE, '%Y-%m-%d').date() <= datetime.datetime.now().date():
                    # 去除無效檢體
                    if row.RESULT_DESC != '無效檢體':
                        data_dict = {'案例編號': row.case_n, '傳染病報告單電腦編號': str(row.REPORT), '姓名(完整)': row.NAME,
                                     '最後一套採檢日期': row.SAMPLE_DATE, '檢體種類': row.SAMPLE_TYPE_DESC, '綜合檢驗結果': row.RESULT_DESC,
                                     '檢驗結果註記': row.COMMENTS}

                        data_list.append(data_dict)
    # print(data_list)
    return analysis(data_list, report_number_list)


def find_latest_date(data_list, _3_date_check_list):
    # 找出最新日期
    final_list = []
    # print(data_list)
    sample_kind_list = ['下呼吸道抽取物', '咽喉擦拭液/鼻咽拭子', '痰液', '唾液(機場專用)']
    for sample in sample_kind_list:
        temp_data = []
        temp_date = []
        for data in data_list:
            if data[4] == sample:
                temp_data.append(data)
                temp_date.append(data[3])

        if temp_date:
            latest_date = max(temp_date)
            for d in temp_data:

                if d[3] == latest_date:
                    if d not in final_list:
                        final_list.append(d)
    val_dict = {}
    for i in final_list:
        if i[4] not in val_dict:
            val_dict[i[4]] = i

        else:
            if type_to_num(i[-2]) > type_to_num(val_dict[i[4]][-2]):
                val_dict[i[4]] = i

    return list(val_dict.values())


def type_to_num(k):

    temp = {'陽性': 3, '陰性': 2, '尚無研判結果': 1, '': 0}

    return temp[k]


# 資料分析
def analysis(data_list, report_number_list):
    final_exam_list = []
    _3_final_list = []

    for rn in report_number_list:
        temp_1 = []
        _3_date_check_list = []
        positive_date_list = []
        for data in data_list:
            if data['傳染病報告單電腦編號'] == str(rn):
                temp = [data['案例編號'], data['傳染病報告單電腦編號'], data['姓名(完整)'], data['最後一套採檢日期'], data['檢體種類'], data['綜合檢驗結果'], data['檢驗結果註記']]
                # print(temp)
                if temp not in temp_1:
                    temp_1.append(temp)
                if data['最後一套採檢日期'] not in _3_date_check_list:
                    _3_date_check_list.append(data['最後一套採檢日期'])
                if data['綜合檢驗結果'] == '陽性':
                    if data['最後一套採檢日期'] not in positive_date_list:
                        positive_date_list.append(data['最後一套採檢日期'])

        temp_1 = sorted(temp_1, key=lambda x: x[3])

        _3_date_check_list = sorted(_3_date_check_list, key=lambda x: datetime.datetime.strptime(x, '%Y-%m-%d').date())

        positive_date_list = sorted(positive_date_list)

        # 檢驗3天用
        if len(_3_date_check_list) >= 3:

            temp_3_final_list = []
            num_date_list = []
            for temp in temp_1:
                if _3_date_check_list[-3] == temp[3] and _3_date_check_list[-3] not in positive_date_list:
                    if temp[-2] == '陰性':
                        if temp not in temp_3_final_list:
                            temp_3_final_list.append(temp)
                            if temp[3] not in num_date_list:
                                num_date_list.append(temp[3])

                if _3_date_check_list[-2] == temp[3] and _3_date_check_list[-2] not in positive_date_list:
                    if temp[-2] == '陰性':
                        if temp not in temp_3_final_list:
                            temp_3_final_list.append(temp)
                            if temp[3] not in num_date_list:
                                num_date_list.append(temp[3])

                if _3_date_check_list[-1] == temp[3] and _3_date_check_list[-1] not in positive_date_list:
                    if temp[-2] == '陰性' or temp[-2] == '尚無研判結果':
                        if temp not in temp_3_final_list:
                            temp_3_final_list.append(temp)
                            if temp[3] not in num_date_list:
                                num_date_list.append(temp[3])

            if len(num_date_list) >= 3:
                _3_final_list += temp_3_final_list

        # 各檢驗種類最新日期
        if temp_1:
            hole_info_list = find_latest_date(temp_1, _3_date_check_list)

        else:
            continue

        if hole_info_list:
            final_exam_list += hole_info_list

    final_exam_list = sorted(final_exam_list, key=lambda x: datetime.datetime.strptime(x[3], '%Y-%m-%d').date())

    final_exam_list = sorted(final_exam_list, key=lambda x: float(x[0][1:]))

    _3_final_list = sorted(_3_final_list, key=lambda x: float(x[0][1:]))

    return [data_inline(final_exam_list, report_number_list), _3_final_list]


def data_inline(data_list, report_number_list):
    # sample_kind_list = ['痰液', '咽喉擦拭液/鼻咽拭子', '下呼吸道抽取物', '唾液(機場專用)']
    final_list = []

    for rn in report_number_list:
        # d_1, s_1, r_1 :痰液  d_2, s_2, r_2 :咽喉擦拭液/鼻咽拭子  d_3, s_3, r_3 :下呼吸道抽取物, d_4, s_4, r_4 :唾液(機場專用)
        model = {'case': '', 'rn': '', 'name': '',
                 'd_1': '無檢驗', 's_1': '無檢驗', 'r_1': '無檢驗', 'c_1': '無資料',
                 'd_2': '無檢驗', 's_2': '無檢驗', 'r_2': '無檢驗', 'c_2': '無資料',
                 'd_3': '無檢驗', 's_3': '無檢驗', 'r_3': '無檢驗', 'c_3': '無資料',
                 'd_4': '無檢驗', 's_4': '無檢驗', 'r_4': '無檢驗', 'c_4': '無資料'}

        for data in data_list:

            if str(rn) == data[1]:
                model['case'] = data[0]
                model['rn'] = data[1]
                model['name'] = data[2]

                if data[4] == '痰液':
                    model['d_1'] = data[3]
                    model['s_1'] = data[4]
                    model['r_1'] = data[5]
                    model['c_1'] = data[6]
                if data[4] == '咽喉擦拭液/鼻咽拭子':
                    model['d_2'] = data[3]
                    model['s_2'] = data[4]
                    model['r_2'] = data[5]
                    model['c_2'] = data[6]
                if data[4] == '下呼吸道抽取物':
                    model['d_3'] = data[3]
                    model['s_3'] = data[4]
                    model['r_3'] = data[5]
                    model['c_3'] = data[6]
                if data[4] == '唾液(機場專用)':
                    model['d_4'] = data[3]
                    model['s_4'] = data[4]
                    model['r_4'] = data[5]
                    model['c_4'] = data[6]
        final_list.append([model['case'], model['rn'], model['name'], model['d_1'], model['s_1'], model['r_1'], model['c_1'],
                           model['d_2'], model['s_2'], model['r_2'], model['c_2'], model['d_3'], model['s_3'], model['r_3'], model['c_3'],
                           model['d_4'], model['s_4'], model['r_4'], model['c_4']
                           ])
        for i in final_list:
            if not i[0]:
                final_list.remove(i)

        final_list = sorted(final_list, key=lambda x: float(x[0][1:]))

    return final_list


def write_xlsx(analyzed_data):
    # 寫成 xlsx
    data_1 = analyzed_data[0]
    data_2 = analyzed_data[1]
    data_len_1 = len(analyzed_data[0])
    data_len_2 = len(analyzed_data[1])

    day = datetime.datetime.now().date()
    hour = datetime.datetime.now().hour
    file_name = ''
    if hour >= 12:
        file_name = f"{day}檢驗結果更新_下午15點版.xlsx"
    if hour < 12:
        file_name = f"{day}檢驗結果更新_上午7點版.xlsx"

    workbook = xlsxwriter.Workbook(os.path.join('.', 'exported_data', 'cdc', file_name))

    worksheet_1 = workbook.add_worksheet('latest')
    worksheet_2 = workbook.add_worksheet('2negative+1_cases')

    table_format = workbook.add_format({'align': 'center',
                                        'valign': 'vcenter',
                                        'size': 12, 'color': 'black'})

    sheet_1_col_param = [{'header': '案例編號', 'format': table_format}, {'header': '傳染病報告單電腦編號', 'format': table_format}, {'header': '姓名(完整)', 'format': table_format},
                         {'header': '採檢日期-痰液', 'format': table_format}, {'header': '檢體種類-痰液', 'format': table_format}, {'header': '檢驗結果-痰液', 'format': table_format},
                         {'header': '檢驗結果註記-痰液', 'format': table_format},
                         {'header': '採檢日期-咽喉擦拭液/鼻咽拭子', 'format': table_format}, {'header': '檢體種類-咽喉擦拭液/鼻咽拭子', 'format': table_format}, {'header': '檢驗結果-咽喉擦拭液/鼻咽拭子', 'format': table_format},
                         {'header': '檢驗結果註記-咽喉擦拭液/鼻咽拭子', 'format': table_format},
                         {'header': '採檢日期-下呼吸道抽取物', 'format': table_format}, {'header': '檢體種類-下呼吸道抽取物', 'format': table_format}, {'header': '檢驗結果-下呼吸道抽取物', 'format': table_format},
                         {'header': '檢驗結果註記-下呼吸道抽取物', 'format': table_format},
                         {'header': '採檢日期-唾液(機場專用)', 'format': table_format}, {'header': '檢體種類-唾液(機場專用)', 'format': table_format}, {'header': '檢驗結果-唾液(機場專用)', 'format': table_format},
                         {'header': '檢驗結果註記-唾液(機場專用)', 'format': table_format},
                         ]

    sheet_2_col_param = [{'header': '案例編號', 'format': table_format}, {'header': '傳染病報告單電腦編號', 'format': table_format},
                         {'header': '姓名(完整)', 'format': table_format}, {'header': '最後一套採檢日期', 'format': table_format},
                         {'header': '檢體種類', 'format': table_format}, {'header': '綜合檢驗結果', 'format': table_format}, {'header': '檢驗結果註記', 'format': table_format}]

    add_table_content_1 = {'data': data_1,
                           'autofilter': True,
                           'columns': sheet_1_col_param}
    add_table_content_2 = {'data': data_2,
                           'autofilter': True,
                           'columns': sheet_2_col_param}

    worksheet_1.add_table(f'A1:S{data_len_1 + 3}', add_table_content_1)
    worksheet_1.set_column(f'A:G', 30)
    worksheet_1.set_column(f'H:K', 40)
    worksheet_1.set_column(f'L:O', 35)
    worksheet_1.set_column(f'P:S', 35)

    worksheet_2.add_table(f'A1:G{data_len_2 + 3}', add_table_content_2)
    worksheet_2.set_column(f'A:G', 30)

    title_format = workbook.add_format({'align': 'center',
                                        'valign': 'vcenter',
                                        'size': 14, 'bold': 4,
                                        'color': 'white',
                                        })

    for i in range(data_len_1 + 1):
        if i == 0:
            worksheet_1.set_row(i, 30, title_format)
        else:
            worksheet_1.set_row(i, 20)

    for i in range(data_len_2 + 1):
        if i == 0:
            worksheet_2.set_row(i, 30, title_format)
        else:
            worksheet_2.set_row(i, 20)

    # 檢驗陽性格式
    positive_format = workbook.add_format({'align': 'center',
                                           'valign': 'vcenter',
                                           'size': 12, 'bold': 4,
                                           'color': 'red',
                                           'bg_color': '#FFC7CE',
                                           'font_color': '#9C0006'
                                           })
    # 無檢驗資料格式
    none_format = workbook.add_format({'align': 'center',
                                       'valign': 'vcenter',
                                       'size': 12,
                                       'color': 'gray',
                                       })

    worksheet_1.conditional_format(f'A1:S{data_len_1 + 3}', {'type': 'text',
                                                             'criteria': 'containing',
                                                             'value': '陽性',
                                                             'format': positive_format})

    worksheet_1.conditional_format(f'A1:S{data_len_1 + 3}', {'type': 'text',
                                                             'criteria': 'containing',
                                                             'value': '無檢驗',
                                                             'format': none_format})

    worksheet_1.conditional_format(f'A1:S{data_len_1 + 3}', {'type': 'text',
                                                             'criteria': 'containing',
                                                             'value': '無資料',
                                                             'format': none_format})

    worksheet_2.conditional_format(f'A1:G{data_len_2 + 3}', {'type': 'text',
                                                             'criteria': 'containing',
                                                             'value': '陽性',
                                                             'format': positive_format})

    worksheet_2.conditional_format(f'A1:G{data_len_2 + 3}', {'type': 'text',
                                                             'criteria': 'containing',
                                                             'value': '無資料',
                                                             'format': none_format})

    workbook.close()

    return file_name


# 下載DB資料
def download_raw_data(file_name, code):
    dsn = cx_Oracle.makedsn('IP', 'port', service_name='DW')

    conn = cx_Oracle.connect(
        user='username',
        password='pwd',
        dsn=dsn,
        encoding='UTF8',
        nencoding='UTF8'
    )

    c = conn.cursor()
    c.execute(code)
    print('code execution completed')

    desc = c.description
    col_name_list = []
    for s in desc:
        col_name_list.append(s[0])

    data_list = c.fetchall()
    conn.close()

    df = pd.DataFrame(data_list, columns=col_name_list)

    if file_name == 'code_1':
        df.to_csv(os.path.join('.', 'imported_data', 'cdc', '2019-nCov_all_exam.csv'), encoding='utf-8')
    if file_name == 'code_2':
        df.to_csv(os.path.join('.', 'imported_data', 'cdc', 'asd.csv'), encoding='utf-8')


# 下載案次號與法傳編號對照表
def download_comparison_data():
    url = "http://IP/share.cgi?ssid=0mUU3e7&fid=0mUU3e7&path=%2F02_%E7%A2%BA%E8%A8%BA%E5%80%8B%E6%A1%88%E7%96%AB%E8%AA%BF&filename=" \
          "%E5%8D%80%E7%AE%A1_%E5%80%8B%E6%A1%88%E6%B3%95%E5%82%B3%E7%B7%A8%E8%99%9F%E5%8F%8A%E6%A1%88%E6%AC%A1%E8%99%9F%E5%B0%8D%E7%85%A7%E8%A1%A8.xlsx&" \
          "openfolder=forcedownload&ep="
    with open(os.path.join('', 'imported_data', 'cdc', '法傳編號與案次號對照表.xlsx'), 'wb') as f:
        f.write(requests.get(url, verify=False).content)


def main():
    # 刪除資料
    delete_xlsx_file_list = glob.glob((os.path.abspath(os.path.join('.', 'exported_data', 'cdc', '*.xlsx'))))
    for file_path in delete_xlsx_file_list:
        if os.path.exists(file_path):
            os.remove(file_path)

    delete_zip_file_list = glob.glob((os.path.abspath(os.path.join('.', 'exported_data', 'cdc', '*.zip'))))
    for file_path in delete_zip_file_list:
        if os.path.exists(file_path):
            os.remove(file_path)

    # 下載檢驗報告資料
    sql_code = {'code_1': f"select "
                          f"ttt2.IDNO,ttt2.NAME,ttt2.REPORT,ttt2.SAMPLE,ttt2.SAMPLE_DATE,ttt2.SAMPLE_TYPE_DESC,ttt2.RESULT_DESC "
                          f"from (select IDNO "
                          f"from CDCDW.DWS_SAMPLE_DETAIL t1 "
                          f"left join CDCDW.USV_INDIVIDUAL_SAS t2 on t1.INDIVIDUAL = t2.INDIVIDUAL "
                          f"where t1.disease in ('19CoV','SICoV','SICV2') "
                          f"and t1.testee_type not in (2,15) "
                          f"and t1.SAMPLE_DATE >= TO_DATE('2020/1/15', 'YYYY/MM/DD') "
                          f"and t1.result = 5 "
                          f"group by IDNO) ttt1 "
                          f"left join (select tt1.IDNO,tt1.NAME,tt2.REPORT,tt1.SAMPLE,tt1.SAMPLE_DATE,tt1.SAMPLE_TYPE_DESC,tt1.RESULT_DESC from "
                          f"(select t2.IDNO,t2.NAME,t1.SAMPLE,t1.SAMPLE_DATE,t3.SAMPLE_TYPE_DESC,t4.RESULT_DESC "
                          f"from CDCDW.DWS_SAMPLE_DETAIL t1 "
                          f"left join CDCDW.USV_INDIVIDUAL_SAS t2 on t1.INDIVIDUAL = t2.INDIVIDUAL "
                          f"left join CDCDW.DIM_SAMPLE_TYPE t3 on t1.SAMPLE_TYPE = t3.SAMPLE_TYPE "
                          f"left join CDCDW.DIM_RESULT t4 on t1.RESULT = t4.RESULT "
                          f"where t1.disease in ('19CoV','SICoV','SICV2') "
                          f"and t1.testee_type not in (2,15) "
                          f"and t1.SAMPLE_DATE >= TO_DATE('2020/1/15', 'YYYY/MM/DD') "
                          f"and t3.SAMPLE_TYPE_DESC in ('下呼吸道抽取物','右肺','鼻咽拭子/咽喉擦拭-病毒','咽喉擦拭液','痰','痰液', '鼻咽擦拭液', '唾液(機場專用)') "
                          f"group by IDNO,NAME,REPORT,SAMPLE,SAMPLE_DATE,SAMPLE_TYPE_DESC,RESULT_DESC )tt1 "
                          f"left join CDCDW.USV_DWS_REPORT_DETAIL_EIC tt2 on tt1.IDNO = tt2.IDNO) ttt2 on ttt1.IDNO = ttt2.IDNO "
                          f"order by NAME, SAMPLE_DATE asc",
                'code_2': f"select t1.SAMPLE, t1.DISEASE,t2.COMMENTS,t2.RESULT_DATE from cdcdw.DWS_SAMPLE_DETAIL t1 left join cdcdw.SAMPLE_PATHOGEN_RESULT t2 on t1.SAMPLE = t2.SAMPLE "
                          f"where t1.DISEASE in ('19CoV','SICV2','SICoV') order by t2.RESULT_DATE desc"}
    # for k, v in sql_code.items():
        # download_raw_data(k, v)

    # 下載法傳編號與按次號對照表
    # download_comparison_data()

    # 執行程式分析
    write_list = data_clean()
    #
    # 寫成 xlsx
    file_path = write_xlsx(write_list)

    e_mail_list = ['wanchin@cdc.gov.tw', 'huang.songen@cdc.gov.tw', 'cpsu@cdc.gov.tw', 'wei-ju@cdc.gov.tw', 'tsungpei@cdc.gov.tw', 'leepinhui@cdc.gov.tw',
                   'pcanita.tw@cdc.gov.tw', 'yclin@cdc.gov.tw', 'yudihsu@cdc.gov.tw', 'liuyl@cdc.gov.tw', 'mengyuchen@cdc.gov.tw',
                   'peiyuanwu@cdc.gov.tw', 'littleka@cdc.gov.tw', 'neo811016@cdc.gov.tw', 'wuhaushing@cdc.gov.tw']

    test_mail = ['yudihsu@cdc.gov.tw']

    if file_path:
        w_file_path = os.path.abspath(os.path.join('.', 'exported_data', 'cdc', file_path))

        zip_name = os.path.abspath(os.path.join('.', 'exported_data', 'cdc', f'{os.path.splitext(os.path.basename(w_file_path))[0]}.zip'))

        password_d = datetime.datetime.strftime(datetime.datetime.now(), '%Y%m%d')

        Popen(['zip', '-j', zip_name, w_file_path, '-P', f'1922{password_d}'])

        while True:
            if os.path.isfile(zip_name):
                break
            time.sleep(0.1)

        mail = SMTP(e_mail_list, [zip_name], os.path.splitext(os.path.basename(zip_name))[0], os.path.splitext(os.path.basename(zip_name))[0])

        # mail.send()


if __name__ == '__main__':
    main()
